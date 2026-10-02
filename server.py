"""Excel-backed reference adapter for the operational attendance report.

The browser receives only report rows and lookup lists.  The adapter can be
replaced by Supabase later without changing the grid interface.
"""

from __future__ import annotations

import json
import mimetypes
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

import openpyxl

PROJECT = Path(__file__).resolve().parent
SOURCE = PROJECT.parent
TIPOLOGIA_FILE = SOURCE / "Tipologia e Indicadores.xlsx"
ATENDIMENTOS_FILE = SOURCE / "Atendimentos.xlsx"
REPORT_HEADERS = ["site", "priority", "quantity", "technician", "base", "failure", "status", "notes", "voltage", "energyOne", "energyTwo"]
TECHNICIAN_BASES = {
    "ADRIANO": "OLI", "ALAN": "CVL", "ANTONIO": "LGP", "BRUNO": "JML", "CLAUDIO": "DIV", "CLEYTON": "LGP",
    "CRISTOVAM": "ITB", "DENILSON": "PNV", "FLAVIO": "BH", "GLEICIANO": "SLG", "HAMILTON": "CLF", "ISAC LIMA": "DIV",
    "JOSE HORTA": "MRN", "JULIANO": "SLG", "LAERCIO": "PNV", "MARCELO": "DIV", "MARCIO": "DIV", "REINALDO": "BH",
    "ROBERTO": "CLF", "SIDNEY": "DIV", "WASHINGTON": "JML", "FERNANDO": "CLF", "APOIO BH": "BH", "TAUA": "SER",
    "APOIO NORTE/LESTE": "MCL", "ANDERSON": "MRN", "MAX CLEI": "ITB",
}
SUPPORT_TECHNICIANS = ["APOIO ZM", "APOIO OESTE", "APOIO BH", "APOIO SUL"]


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, timedelta):
        seconds = int(value.total_seconds())
        return f"{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


def site_key(value: object) -> str:
    return re.sub(r"\s+", "", text(value).upper())


def technician_key(value: object) -> str:
    normalized = unicodedata.normalize("NFD", text(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]", "", normalized.upper())


def distinct(values: list[str]) -> list[str]:
    return list(dict.fromkeys(item for item in values if item))


class Catalog:
    def __init__(self) -> None:
        self.sites: dict[str, dict[str, str]] = {}
        self.technicians: list[str] = list(TECHNICIAN_BASES)
        self.technician_bases: dict[str, str] = {technician_key(name): base for name, base in TECHNICIAN_BASES.items()}
        self.failures: list[str] = [
            "FALHA DE DISJUNTOR",
            "GMG OPERANDO",
            "FALHA NO GMG",
            "ALTA TEMPERATURA",
            "TROCAR DE CALOR",
            "INOPERANTE",
            "BAIXO NÍVEL COMBUSTÍVEL",
            "BATERRY CURRENT OUT OF RANGE",
            "FALHA INVERSOR",
            "FALHA DE AC",
            "BATERIA EM DESCARGA",
        ]
        self.statuses: list[str] = [
            "ACIONADO",
            "EM DESLOCAMENTO",
            "GMG ACOPLADO",
            "MONITORANDO",
            "INDISP. TÉCNICA",
            "SEM ACESSO",
        ]
        self.report_rows: list[dict[str, str]] = []
        self.attendimentos_available = False
        self.load_tipologia()
        self.load_atendimentos()
        # Operational support options complement the workbook's technician list.
        existing = {technician_key(name) for name in self.technicians}
        self.technicians.extend(name for name in SUPPORT_TECHNICIANS if technician_key(name) not in existing)

    def load_tipologia(self) -> None:
        workbook = openpyxl.load_workbook(TIPOLOGIA_FILE, read_only=True, data_only=True)
        sheet = workbook["Tipologia"]
        for row in sheet.iter_rows(min_row=6, values_only=True):
            code = site_key(row[3] if len(row) > 3 else None)
            if not code:
                continue
            self.sites[code] = {
                "site": code,
                "priority": text(row[6] if len(row) > 6 else None),
                "quantity": text(row[11] if len(row) > 11 else None),
                "base": text(row[28] if len(row) > 28 else None),
                "energyI": text(row[30] if len(row) > 30 else None),
                "energyII": text(row[31] if len(row) > 31 else None),
            }
        contacts = workbook["Contatos"]
        technicians: list[str] = []
        for row in contacts.iter_rows(min_row=2, values_only=True):
            name = text(row[0] if row else None)
            cluster = text(row[1] if len(row) > 1 else None).upper()
            if name and cluster in {"BH", "METROPOLITANO"}:
                technicians.append(name)
        if not self.technicians:
            self.technicians = distinct(technicians)
        workbook.close()

    def load_atendimentos(self) -> None:
        """Extract the visible SITES grid and its validation lists when the workbook is free.

        Excel can hold the source open during a supervisor's shift.  In that
        case Tipologia-based lookup still works, and this method is retried on
        a later server start without fabricating dropdown values.
        """
        try:
            workbook = openpyxl.load_workbook(ATENDIMENTOS_FILE, data_only=True)
        except PermissionError:
            return
        except OSError:
            return
        try:
            sheet = workbook["SITES"]
            header_row = next((row for row in range(1, min(sheet.max_row, 12) + 1) if site_key(sheet.cell(row, 1).value) in {"ESTAÇÃO", "ESTACAO"}), None)
            if not header_row:
                return
            for values in sheet.iter_rows(min_row=header_row + 1, max_col=11, values_only=True):
                row = {key: text(values[index] if len(values) > index else None) for index, key in enumerate(REPORT_HEADERS)}
                if not any(row.values()):
                    continue
                self.report_rows.append(row)
            self.technicians = distinct(self.technicians + [row["technician"] for row in self.report_rows])
            self.technicians.sort(key=str.casefold)
            self.failures = distinct(self.failures + [row["failure"] for row in self.report_rows])
            self.statuses = distinct(self.statuses + [row["status"] for row in self.report_rows])
            technicians_sheet = workbook["TÉCNICOS"]
            official_technicians: list[str] = []
            for values in technicians_sheet.iter_rows(min_row=5, max_col=2, values_only=True):
                name, base = (text(values[index] if len(values) > index else None) for index in (0, 1))
                if not name:
                    continue
                official_technicians.append(name)
                self.technician_bases[technician_key(name)] = base
            if official_technicians:
                self.technicians = official_technicians
            self.read_validation_lists(workbook, sheet)
            self.attendimentos_available = True
        finally:
            workbook.close()

    def read_validation_lists(self, workbook: openpyxl.Workbook, sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
        """Read list validation source ranges and merge them into each matching column."""
        targets = {4: "technicians", 6: "failures", 7: "statuses"}
        for validation in sheet.data_validations.dataValidation:
            if validation.type != "list" or not validation.formula1:
                continue
            covered_columns = {cell.column for cell_range in validation.sqref.ranges for row in sheet[cell_range.coord] for cell in row}
            matching = [name for column, name in targets.items() if column in covered_columns]
            if not matching:
                continue
            values = self.validation_values(workbook, validation.formula1)
            for name in matching:
                setattr(self, name, distinct(getattr(self, name) + values))

    def validation_values(self, workbook: openpyxl.Workbook, formula: str) -> list[str]:
        source = formula.lstrip("=")
        defined = workbook.defined_names.get(source)
        destinations = list(defined.destinations) if defined else []
        if "!" in source and not destinations:
            title, cells = source.split("!", 1)
            destinations = [(title.strip("'"), cells)]
        values: list[str] = []
        for title, cells in destinations:
            if title not in workbook.sheetnames:
                continue
            for row in workbook[title][cells.replace("$", "")]:
                for cell in row if isinstance(row, tuple) else (row,):
                    values.append(text(cell.value))
        return values

    def search_sites(self, query: str) -> list[str]:
        needle = site_key(query)
        return [code for code in self.sites if needle in code][:16]


CATALOG = Catalog()


class ReportHandler(SimpleHTTPRequestHandler):
    def send_json(self, value: object, status: HTTPStatus = HTTPStatus.OK) -> None:
        content = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path == "/api/health":
            self.send_json({"source": TIPOLOGIA_FILE.name, "sites": len(CATALOG.sites), "reportSourceAvailable": CATALOG.attendimentos_available})
            return
        if parsed.path == "/api/options":
            self.send_json({"technicians": CATALOG.technicians, "technicianBases": CATALOG.technician_bases, "failures": CATALOG.failures, "statuses": CATALOG.statuses})
            return
        if parsed.path == "/api/report":
            self.send_json({"rows": CATALOG.report_rows})
            return
        if parsed.path == "/api/sites":
            self.send_json({"items": CATALOG.search_sites(parse_qs(parsed.query).get("q", [""])[0])})
            return
        if parsed.path.startswith("/api/sites/"):
            record = CATALOG.sites.get(site_key(unquote(parsed.path.rsplit("/", 1)[-1])))
            self.send_json(record or {"message": "Site não localizado na Tipologia."}, HTTPStatus.OK if record else HTTPStatus.NOT_FOUND)
            return
        super().do_GET()

    def translate_path(self, path: str) -> str:
        relative = "index.html" if urlparse(path).path in {"", "/"} else urlparse(path).path.lstrip("/")
        target = (PROJECT / relative).resolve()
        return str(target if PROJECT in target.parents or target == PROJECT else PROJECT / "index.html")

    def end_headers(self) -> None:
        self.send_header("Cache-Control", "no-store")
        super().end_headers()


if __name__ == "__main__":
    server = ThreadingHTTPServer(("127.0.0.1", 8080), ReportHandler)
    print("Relatório disponível em http://127.0.0.1:8080")
    server.serve_forever()
